"""The 🧮 Excel Builder, rebuilt to look and behave like Excel's own PivotTable (28 Oct 2026).

Manav: *"the broad level idea is to mimic the exact functionality of a pivot table we use on excel"* — and,
asked how far: *"a rebuild that looks and behaves like Excel's own pane"*.

What is Excel here and what is not:

  THE SCREEN IS EXCEL'S. A ribbon (PivotTable Analyze · Design), a sheet with column letters and row numbers,
  report filters above the table, "Row Labels ▾" / "Column Labels ▾", ⊞/⊟ on every group, the PivotTable
  Fields pane with drag and drop between FILTERS · COLUMNS · ROWS · Σ VALUES, the chip menus (Move Up …
  Remove Field), Value Field Settings, Number Format, compact/outline/tabular, subtotals top/bottom/off,
  grand totals on/off by rows and columns, repeat labels, blank rows, banded rows, cell selection with
  Sum/Count/Average in the status bar, Defer Layout Update. All of that lives in `pivot_pane.js`.

  THE ARITHMETIC IS THE DASHBOARD'S. Every number comes from `pivotx.build` — the engine the 13 Oct builder
  used, with every subtotal and total re-derived from the rows it covers (a ratio's total is a true ratio;
  see [[feedback-aggregate-ratios-in-pairs]]). This module only translates the pane's layout into an engine
  request and hands the result to the screen.

★ ONE CALCULATION, MANY VIEWS. The engine is always asked for the FULL tree — subtotals at every level, both
grand totals — so switching layout, moving subtotals, turning grand totals off or collapsing a group is
instant in the browser. Only a change to WHAT is measured (fields, values, filters, dates, sort, top N) comes
back here.

★ THE EXCEL FILE IS THE SCREEN. "Export" sends the sheet exactly as drawn — layout, names, number formats,
collapsed groups — and `to_excel` writes that, so the file never disagrees with what was on screen.

★ NO CUSTOMER MOBILES, as a field, an item or a source row (inherited from `pivotx`).
"""
from __future__ import annotations

import base64
import io
import json
from pathlib import Path

import pandas as pd
import streamlit as st

import pivotx as P
import pivot_ui as PU

HERE = Path(__file__).resolve().parent
ITEM_CAP = 2500                     # items offered in a filter list (Style code alone runs to thousands)

# Excel's wording on screen ↔ the engine's. The pane stores the engine's name, the screen shows Excel's.
SHOW_AS_EXCEL = {
    "Value": "No Calculation",
    "% of grand total": "% of Grand Total",
    "% of column total": "% of Column Total",
    "% of row total": "% of Row Total",
    "% of parent row": "% of Parent Row Total",
    "Change vs previous column": "Difference From",
    "% change vs previous column": "% Difference From",
    "Running total across": "Running Total In",
}

# Excel's number formats for the file. "auto" = the dashboard's own per kind (Indian grouping).
_XNUM = {
    "number": lambda d: "##,##0" + ("." + "0" * d if d else ""),
    "indian": lambda d: '[>=10000000]##\\,##\\,##\\,##0' + ("." + "0" * d if d else "") + ';[>=100000]##\\,##\\,##0'
                        + ("." + "0" * d if d else "") + ';##,##0' + ("." + "0" * d if d else ""),
    "currency": lambda d: '"₹"[>=10000000]##\\,##\\,##\\,##0' + ("." + "0" * d if d else "") + ';"₹"[>=100000]##\\,##\\,##0'
                          + ("." + "0" * d if d else "") + ';"₹"##,##0' + ("." + "0" * d if d else ""),
    "lakh": lambda d: "0" + ("." + "0" * d if d else "") + '" L"',
    "crore": lambda d: "0" + ("." + "0" * d if d else "") + '" Cr"',
    "percent": lambda d: "0" + ("." + "0" * d if d else "") + "%",
}


# =========================================================================== #
#  The layout ("ui") — what the pane edits — and its translation to the engine
# =========================================================================== #
def default_ui(feed: str) -> dict:
    return from_preset(feed, PU.PRESETS[feed][PU.DEFAULT[feed]])


def from_preset(feed: str, p: dict) -> dict:
    """A 13-Oct preset (or saved layout) as a pane layout."""
    M = P.measures(feed)
    vals = [{"m": m, "how": h, "name": "", "fmt": {"cat": "auto", "dec": 0}}
            for m, h in (tuple(v) for v in p.get("values", [])) if m in M]
    sb = p.get("sort_by", "Label")
    sort = ({"by": "label", "desc": False} if sb in ("Label", "Label (A to Z)") else
            {"by": "label", "desc": True} if sb == "Label (Z to A)" else
            {"by": sb, "desc": bool(p.get("descending", True))})
    flt = p.get("filters") or {}
    return {
        "rows": list(p.get("rows", [])), "cols": list(p.get("cols", [])), "filters_area": list(flt.keys()),
        "sel": {f: {"mode": v.get("mode", "include"), "items": list(v.get("items", []))} for f, v in flt.items()},
        "values": vals,
        "period": p.get("period") or "FY to date · this year and last", "from": None, "to": None,
        "sort": sort, "top": int(p.get("top", 0) or 0),
        # display only — the browser applies these without recalculating
        "layout": {"Compact": "compact", "Tabular": "tabular"}.get(p.get("layout", "Compact"), "compact"),
        "subtotals": "top" if p.get("subtotals", True) else "off",
        "grand_rows": bool(p.get("grand_rows", True)), "grand_cols": bool(p.get("grand_cols", True)),
        "repeat": False, "blank": False, "banded_rows": False, "banded_cols": False,
        "row_headers": True, "col_headers": True, "buttons": True, "field_headers": True, "style": "light16",
        "names": {},
    }


def _ranges(ui: dict, asof) -> tuple[list, str]:
    per = P.periods(asof)
    if ui.get("period") == "Custom" and ui.get("from") and ui.get("to"):
        return [[str(ui["from"]), str(ui["to"])]], "Custom"
    name = ui.get("period") if ui.get("period") in per else "FY to date · this year and last"
    return [[str(pd.Timestamp(a).date()), str(pd.Timestamp(b).date())] for a, b in per[name]], name


def to_spec(ui: dict, feed: str, asof) -> P.Spec:
    """The engine request: ALWAYS the full tree (subtotals everywhere, both grand totals, tabular), so the
    browser can show any layout of it without asking again."""
    M = P.measures(feed)
    used = set(ui["rows"]) | set(ui["cols"]) | set(ui.get("filters_area", []))
    filters = {f: {"mode": s.get("mode", "include"), "items": list(s.get("items", []))}
               for f, s in (ui.get("sel") or {}).items() if f in used and s.get("items")}
    ranges, period = _ranges(ui, asof)
    s = ui.get("sort") or {}
    sort_by = (("Label (Z to A)" if s.get("desc") else "Label") if s.get("by", "label") == "label"
               else s["by"] if s.get("by") in M else "Label")
    return P.Spec(rows=list(ui["rows"]), cols=list(ui["cols"]),
                  values=[(v["m"], v.get("how", "Value")) for v in ui["values"] if v["m"] in M],
                  filters=filters, ranges=ranges, period=period, layout="Tabular", subtotals=True,
                  grand_rows=True, grand_cols=True, sort_by=sort_by, descending=bool(s.get("desc", True)),
                  top=int(ui.get("top", 0) or 0))


def _clean(x):
    if x is None:
        return None
    try:
        f = float(x)
    except (TypeError, ValueError):
        return None
    return None if f != f else f


def grid_payload(g: dict, feed: str) -> dict:
    """The engine's grid, compact enough to ship to the browser."""
    M = P.measures(feed)
    return {
        "rf": g["row_fields"], "cf": g["col_fields"], "heads": g["heads"],
        "colent": [[k, list(ck)] for k, ck in g["colent"]],
        "nv": g["nvals"], "kinds": g["kinds"][: g["nvals"]] if g["colent"] else g["kinds"],
        "vkinds": [M[m].kind for (_c, _k, m, _h) in g["leaf"][: g["nvals"]]],
        "rows": [[r["kind"], list(r["key"]), [_clean(x) for x in r["values"]]] for r in g["rows"]],
        "rows_used": g["rows_used"], "rows_in": g["rows_in"], "span": g["span"],
        "applied": g["filters"], "hidden": g["hidden_items"],
    }


@st.cache_data(show_spinner=False, max_entries=12)
def _built(_w, feed: str, spec_json: str, token):
    spec = P.Spec(**json.loads(spec_json))
    spec.values = [tuple(v) for v in spec.values]
    try:
        g = P.build(_w, feed, spec, prepared=True)
    except P.TooBig as e:
        return {"error": str(e)}
    except Exception as e:                                  # surface on the sheet, never crash the tab
        return {"error": f"Could not build that table: {e}"}
    if g.get("empty"):
        return {"error": g["note"]}
    return {"grid": grid_payload(g, feed)}


def _spec_json(spec: P.Spec) -> str:
    d = dict(spec.__dict__)
    d["values"] = [list(v) for v in spec.values]
    return json.dumps(d, sort_keys=True, default=str)


def notes(ui: dict, feed: str, asof, grid: dict | None) -> list[str]:
    """The message bar: what a reader must know to read the table right (carried over from 13 Oct)."""
    out = []
    hows = [v.get("how", "Value") for v in ui["values"]]
    growth = any(h.startswith("% change") or h.startswith("Change") for h in hows)
    finer = {"FY Half", "FY Quarter", "Fiscal Month", "Month", "Quarter", "Week", "Day", "Day of month"}
    if growth and "Financial Year" in ui["cols"] and ui.get("period") in ("All data", "Custom"):
        if finer & set(ui["rows"]):
            out.append(f"The latest period is only part-way through (data to {pd.Timestamp(asof):%d %b %Y}), "
                       "so its change compares part of it with all of last year's.")
        else:
            out.append("This compares a part year with a whole year: this FY so far against all of last FY, so the "
                       "change reads low. Choose 'FY to date · this year and last' to compare the same dates.")
    if growth:
        out.append("Growth here counts every store in the period, including new and closed ones; the GD sheets "
                   "use like to like.")
    if feed == "vfl" and any(v["m"] == "Discount %" for v in ui["values"]):
        out.append("Discount % is a flag, not a depth: since June the promotion amount equals the bill amount "
                   "when a promotion is recorded.")
    if grid and grid.get("hidden"):
        out.append(f"Top {ui.get('top')} kept in each group; {grid['hidden']:,} smaller items hidden "
                   "(totals still include them).")
    return out


# =========================================================================== #
#  The Excel file — written from the sheet exactly as the screen drew it
# =========================================================================== #
STYLES = {   # header fill, header rule, header text, subtotal fill, band fill, grand fill, grand rule
    "light16": ("DDEBF7", "9BC2E6", "000000", None, "EEF5FB", "DDEBF7", "9BC2E6"),
    "light9": ("FFFFFF", "A6A6A6", "000000", None, "F2F2F2", "FFFFFF", "000000"),
    "medium2": ("4472C4", "4472C4", "FFFFFF", "D9E1F2", "D9E1F2", "B4C6E7", "4472C4"),
    "peanuts": ("7A1F2B", "7A1F2B", "FFFFFF", "F6EFE4", "FDFBF8", "EFE3D0", "7A1F2B"),
}


def to_excel(sheet: dict, *, title: str, about: dict, source: pd.DataFrame | None = None) -> bytes:
    """`sheet` is what the browser drew: {"cols": [widths], "head": [[cell]], "body": [{"kind", "level",
    "hidden", "cells": [cell]}], "filters": [[name, shown]], "style": key, "freeze": [row, col]}; a cell is
    {"t": text} or {"v": number, "nf": {"cat", "dec"}, "k": kind} with "b" (bold) and "i" (indent)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    hf, hr, ht, sf, bf, gf, gr = STYLES.get(sheet.get("style"), STYLES["light16"])
    wb = Workbook()
    ws = wb.active
    ws.title = "PivotTable"
    font = "Calibri"
    r = 1
    for name, shown in sheet.get("filters", []):
        ws.cell(row=r, column=1, value=name).font = Font(name=font, bold=True)
        ws.cell(row=r, column=2, value=shown).font = Font(name=font)
        r += 1
    r += 1 if sheet.get("filters") else 2           # as on screen: filters, a blank row, then the pivot (row 3)
    top = r
    for row in sheet.get("head", []):
        for j, c in enumerate(row, start=1):
            cell = ws.cell(row=r, column=j, value=c.get("t") or None)
            cell.font = Font(name=font, bold=True, color=ht)
            if c.get("fill", True) and sheet.get("col_headers", True):
                cell.fill = PatternFill("solid", fgColor=hf)
            cell.border = Border(bottom=Side(style="thin", color=hr))
            cell.alignment = Alignment(horizontal=c.get("al", "left"))
        r += 1
    body0 = r
    for n, row in enumerate(sheet.get("body", [])):
        kind = row.get("kind", "item")
        fill = (gf if kind == "grand" else sf if kind in ("subtotal", "group") and sf else
                bf if (sheet.get("banded_rows") and n % 2) else None)
        for j, c in enumerate(row["cells"], start=1):
            if "v" in c:
                v = c.get("v")
                cat, dec = (c.get("nf") or {}).get("cat", "auto"), int((c.get("nf") or {}).get("dec", 0))
                if v is not None and cat == "lakh":
                    v = v / 1e5
                elif v is not None and cat == "crore":
                    v = v / 1e7
                cell = ws.cell(row=r, column=j, value=v)
                cell.number_format = (P._XFMT.get(c.get("k", ""), "General") if cat == "auto"
                                      else _XNUM.get(cat, _XNUM["number"])(dec))
                cell.alignment = Alignment(horizontal="right")
                neg = v is not None and v < 0 and str(c.get("k", "")).endswith("chg")
                cell.font = Font(name=font, bold=bool(c.get("b")), color=("C00000" if neg else "000000"))
            else:
                cell = ws.cell(row=r, column=j, value=c.get("t") or None)
                cell.font = Font(name=font, bold=bool(c.get("b")))
                cell.alignment = Alignment(horizontal="left", indent=int(c.get("i", 0) or 0))
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
            if kind == "grand":
                cell.border = Border(top=Side(style="thin", color=gr), bottom=Side(style="double", color=gr))
        lvl = int(row.get("level", 0) or 0)
        if lvl:
            ws.row_dimensions[r].outline_level = min(7, lvl)
        if row.get("hidden"):
            ws.row_dimensions[r].hidden = True
        r += 1
    ws.sheet_properties.outlinePr.summaryBelow = bool(sheet.get("summary_below", False))
    for j, w in enumerate(sheet.get("cols", []), start=1):
        ws.column_dimensions[get_column_letter(j)].width = max(8, min(60, float(w) / 7.0))
    fr, fc = sheet.get("freeze", [0, 0])
    if fr or fc:
        ws.freeze_panes = ws.cell(row=top + int(fr), column=1 + int(fc))
    ws.sheet_view.showGridLines = True

    if source is not None and len(source):
        sh = wb.create_sheet("Source rows")
        src = source.drop(columns=[c for c in source.columns if c in P._SOURCE_DROP or c.startswith("_f")], errors="ignore")
        src = src.loc[:, ~src.columns.duplicated()]
        cols = list(src.columns)
        sh.append([str(c) for c in cols])
        for rec in src.itertuples(index=False):
            sh.append([(x.to_pydatetime() if isinstance(x, pd.Timestamp) else
                        (None if (isinstance(x, float) and pd.isna(x)) else x)) for x in rec])
        t = Table(displayName="SourceRows", ref=f"A1:{get_column_letter(len(cols))}{len(src) + 1}")
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        sh.add_table(t)
        sh.freeze_panes = "A2"
        for j, c in enumerate(cols, start=1):
            sh.column_dimensions[get_column_letter(j)].width = min(max(len(str(c)) + 2, 10), 28)

    ab = wb.create_sheet("How this was built")
    ab.cell(row=1, column=1, value=title).font = Font(name=font, bold=True, size=13)
    for i, (k, val) in enumerate(about.items(), start=3):
        ab.cell(row=i, column=1, value=k).font = Font(name=font, bold=True)
        ab.cell(row=i, column=2, value=str(val)).alignment = Alignment(wrap_text=True, vertical="top")
    ab.column_dimensions["A"].width = 24
    ab.column_dimensions["B"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# =========================================================================== #
#  The tab
# =========================================================================== #
_COMP = None


def _component():
    global _COMP
    if _COMP is None:
        _COMP = st.components.v2.component(
            "peanuts_pivot_pane",
            html='<div class="pv-root"></div>',
            css=(HERE / "assets" / "pivot_pane.css").read_text(encoding="utf-8"),
            js=(HERE / "assets" / "pivot_pane.js").read_text(encoding="utf-8"),
        )
    return _COMP


def _k(key, n):
    return f"{key}__{n}"


def _ui_param(ui: dict) -> str:
    return base64.urlsafe_b64encode(json.dumps(ui, separators=(",", ":"), default=str).encode()).decode().rstrip("=")


def _ui_from_param(s: str) -> dict | None:
    try:
        d = json.loads(base64.urlsafe_b64decode(s + "=" * (-len(s) % 4)).decode())
        return d if isinstance(d, dict) and "rows" in d and "values" in d else None
    except Exception:
        return None


def _valid(ui: dict, feed: str) -> dict:
    """Drop anything the feed no longer has (a field renamed, a measure retired) instead of failing."""
    F, M = set(P.field_names(feed)), P.measures(feed)
    base = default_ui(feed)
    out = {**base, **{k: v for k, v in ui.items() if k in base}}
    out["rows"] = [f for f in out["rows"] if f in F]
    out["cols"] = [f for f in out["cols"] if f in F and f not in out["rows"]]
    out["filters_area"] = [f for f in out.get("filters_area", []) if f in F and f not in out["rows"] + out["cols"]]
    out["values"] = [v for v in out["values"] if isinstance(v, dict) and v.get("m") in M
                     and v.get("how", "Value") in P.show_as_for(feed, v["m"])]
    return out


def render(df: pd.DataFrame, *, feed: str, key: str, asof):
    """★ A FRAGMENT (28 Oct 2026): a change in the pivot re-runs THIS tab only, not the whole dashboard — so the
    sheet answers in a second instead of greying out while every other report reloads."""
    st.fragment(_render)(df, feed=feed, key=key, asof=asof)


def _render(df: pd.DataFrame, *, feed: str, key: str, asof):
    ss = st.session_state
    token = (feed, len(df), str(pd.to_datetime(df["date"]).max()))
    W = PU._prepared(df, feed, token)
    M = P.measures(feed)
    uk, rk = _k(key, "ui"), _k(key, "rev")

    if uk not in ss:
        qp = st.query_params.get(f"px_{feed}")
        parsed = _ui_from_param(qp) if qp else None          # a bookmarked or shared table opens as it was
        ss[uk] = _valid(parsed, feed) if parsed else default_ui(feed)
        ss[rk] = 0

    comp = _component()

    def on_change():
        res = ss.get(f"{key}_cmp")
        ev = res.get("event") if hasattr(res, "get") else getattr(res, "event", None)
        if not ev:
            return
        typ = ev.get("type")
        if typ == "ui" and isinstance(ev.get("ui"), dict):
            ss[uk] = _valid(ev["ui"], feed)
            ss[rk] = max(int(ev.get("rev", 0)), ss.get(rk, 0))
        elif typ == "export":
            ss[_k(key, "export")] = ev
        elif typ == "save":
            saved = dict(ss.get(f"px_saved_{feed}") or {})
            if ev.get("name"):
                saved[str(ev["name"])[:60]] = ss[uk]
            ss[f"px_saved_{feed}"] = saved

    ui = ss[uk]
    try:
        st.query_params[f"px_{feed}"] = _ui_param(ui)
    except Exception:
        pass

    # the calculation (cached on the request, so a re-render with the same layout costs nothing)
    if ui["values"]:
        spec = to_spec(ui, feed, asof)
        out = _built(W, feed, _spec_json(spec), token)
    else:
        spec, out = None, {"error": "Choose a field for Σ VALUES — tick a measure in the field list or drag one "
                                    "into the Values box."}

    # items for every field that has a filter list on screen (Filters area, row and column labels)
    items = {}
    for f in dict.fromkeys(ui["rows"] + ui["cols"] + ui.get("filters_area", [])):
        try:
            its = PU._field_items(W, feed, f, token)
        except Exception:
            its = []
        items[f] = {"items": [str(x) for x in its[:ITEM_CAP]], "more": max(0, len(its) - ITEM_CAP)}

    # an export the browser asked for on the previous run: write it now and hand it back
    download = None
    ex = ss.pop(_k(key, "export"), None)
    if ex and isinstance(ex.get("sheet"), dict) and spec is not None:
        src = None
        if ex.get("source"):
            src = P.source_rows(df, feed, spec)
            if len(src) > 250000:
                src = None
        title = f"{'VFL' if feed == 'vfl' else 'Portfolio'} pivot"
        about = {
            "Feed": "VFL (one row per bill line)" if feed == "vfl" else "Portfolio (one row per store per day)",
            "Rows": ", ".join(ui["rows"]) or "(none)", "Columns": ", ".join(ui["cols"]) or "(none)",
            "Values": "; ".join((v.get("name") or v["m"]) + ("" if v.get("how", "Value") == "Value"
                                                              else f" — {SHOW_AS_EXCEL[v['how']]}") for v in ui["values"]),
            "Period": f"{spec.period}: " + " + ".join(f"{a} to {b}" for a, b in spec.ranges) if spec.ranges else spec.period,
            "Filters": "; ".join((out.get("grid") or {}).get("applied", [])) or "(none)",
            "Layout": f"{ui['layout']} form; subtotals {ui['subtotals']}",
            "Source rows used": f"{(out.get('grid') or {}).get('rows_used', 0):,}",
            "Built": f"{pd.Timestamp.now():%d %b %Y %H:%M}",
            "How totals work": "Every subtotal and total is worked out again from the rows it covers, so a ratio's "
                               "total (ATV, UPT, ASP, conversion) is a true ratio, not the sum or average of the "
                               "ratios above it.",
            "Lakhs and crores": "Values formatted in lakhs or crores are written divided by 1,00,000 or 1,00,00,000 "
                                "(Excel's number formats cannot scale by those).",
            "Customer mobiles": "Never included. Customers can be counted, not listed.",
        }
        try:
            data = to_excel(ex["sheet"], title=title, about=about, source=src)
            download = {"id": ex.get("id"), "name": f"{title} {pd.Timestamp(asof):%d-%m-%Y}.xlsx",
                        "b64": base64.b64encode(data).decode()}
        except Exception as e:
            download = {"id": ex.get("id"), "error": f"Could not write the Excel file: {e}"}

    fields = [{"name": f, "group": g} for g, fl in P.GROUPS[feed] for f in fl]
    measures = [{"name": m, "kind": M[m].kind, "additive": M[m].additive, "parts": list(M[m].parts),
                 "show": P.show_as_for(feed, m)} for m in M]
    presets = {n: from_preset(feed, p) for n, p in PU.PRESETS[feed].items()}
    saved = {n: v for n, v in (ss.get(f"px_saved_{feed}") or {}).items()}
    pnames = list(P.periods(asof)) + ["Custom"]
    dmin, dmax = pd.to_datetime(W["date"]).min(), pd.to_datetime(W["date"]).max()

    payload = {
        "feed": feed, "ui": ui, "rev": ss.get(rk, 0), "fields": fields, "measures": measures,
        "show_names": SHOW_AS_EXCEL, "presets": presets, "saved": saved, "periods": pnames,
        "dmin": f"{dmin:%Y-%m-%d}", "dmax": f"{dmax:%Y-%m-%d}", "items": items,
        "grid": out.get("grid"), "error": out.get("error"),
        "notes": notes(ui, feed, asof, out.get("grid")) if out.get("grid") else [],
        "period_dates": (" + ".join(f"{pd.Timestamp(a):%d %b %y} – {pd.Timestamp(b):%d %b %y}" for a, b in spec.ranges)
                         if spec is not None and spec.ranges else "all data"),
        "download": download,
    }
    comp(key=f"{key}_cmp", data=payload, on_event_change=on_change, height="content")
