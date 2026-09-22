"""The Excel builder — pivots off the raw feed, for whoever needs them.

Frames are built inside the test so CI verifies the rules without the feed —
see [[feedback-test-where-it-runs]].
"""
import pandas as pd
import pytest

import pivot as PV


def _vfl():
    import loader as L
    # ★ DELIBERATELY ASYMMETRIC: Agartala writes two bills a day against
    # Jayanagar's one, so a weighted ATV cannot coincide with the plain mean
    # of the two — which is the whole thing being tested.
    rows = []
    for i, d in enumerate(pd.date_range("2026-04-01", "2026-05-31")):
        for store, amt, n in (("Agartala", 1000.0, 2), ("Jayanagar", 2000.0, 1)):
          for b in range(n):
            rows.append({
                "date": d, L.COL_STORE_LABEL: store, L.COL_AMOUNT: amt,
                "net_amount": amt, L.COL_UNITS: 2.0, L.COL_PROMO: 0.0,
                "mobile_clean": f"m{i%7}", "bill_uid": f"{store}-{d:%Y%m%d}-{b}",
                L.COL_REGION: "East & NE", L.COL_BRAND: "Manyavar",
                L.COL_DIVISION: "KURTA SET", L.COL_MWC: "MEN",
                L.COL_SECTION: "S1", L.COL_DEPARTMENT: "D1", L.COL_SIZE: "M",
                L.COL_COLOR_NAME: "TEAL", L.COL_COLOR: "327-Teal",
                L.COL_STYLE: "SDES503", L.COL_SALESPERSON: "ASHIS DAS",
            })
    return pd.DataFrame(rows)


def _pf():
    rows = []
    for d in pd.date_range("2026-04-01", "2026-05-31"):
        for brand, loc, code, amt in (("MANYAVAR", "City Centre", 56, 900.0),
                                      ("TURTLE", "City Centre", 98, 100.0)):
            rows.append({"date": d, "brand": brand, "location": loc,
                         "city": "SILIGURI", "region": "East & NE",
                         "code": code, "sales": amt, "BILL": "3",
                         "QTY": "5", "FOOTFALL": "", "Day Target": "1,000"})
    return pd.DataFrame(rows)


# --------------------------------------------------------------------------- #
# What the feeds offer
# --------------------------------------------------------------------------- #
def test_customer_mobiles_are_never_a_dimension():
    """The feed carries them and this repo is public. Customers are countable,
    never listable."""
    f = PV.fields("vfl")
    every = " ".join(f["time"] + f["cats"] + f["measures"]).lower()
    assert "mobile" not in every
    assert "Unique Customers" in f["measures"]      # counted, not listed


def test_both_feeds_offer_dimensions_and_measures():
    for feed in ("vfl", "portfolio"):
        f = PV.fields(feed)
        assert f["time"] and f["cats"] and f["measures"]


# --------------------------------------------------------------------------- #
# The arithmetic
# --------------------------------------------------------------------------- #
def test_the_total_ties_to_the_feed():
    """A pivot that quietly drops rows is worse than no pivot."""
    df = _vfl()
    out, meta = PV.build(df, feed="vfl", rows=["Store"], measures=["Sales (₹)"])
    import loader as L
    assert meta["grand"]["Sales (₹)"] == pytest.approx(df[L.COL_AMOUNT].sum())
    assert out["Sales (₹)"].sum() == pytest.approx(df[L.COL_AMOUNT].sum())


def test_a_ratio_total_is_derived_not_averaged():
    """ATV over two stores is their total sales over their total bills, not
    the mean of two ATVs. See [[feedback-aggregate-ratios-in-pairs]]."""
    df = _vfl()
    out, meta = PV.build(df, feed="vfl", rows=["Store"],
                         measures=["Sales (₹)", "Bills", "Avg Bill Value / ATV (₹)"])
    derived = meta["grand"]["Sales (₹)"] / meta["grand"]["Bills"]
    assert meta["grand"]["Avg Bill Value / ATV (₹)"] == pytest.approx(derived)
    assert meta["grand"]["Avg Bill Value / ATV (₹)"] != pytest.approx(
        out["Avg Bill Value / ATV (₹)"].mean())


def test_a_pivoted_ratio_column_is_derived_at_its_own_level():
    df = _vfl()
    out, meta = PV.build(df, feed="vfl", rows=["Store"], cols="Month",
                         measures=["Sales (₹)", "Avg Bill Value / ATV (₹)"])
    col = "Apr 2026 · Avg Bill Value / ATV (₹)"
    apr = df[df["date"] < pd.Timestamp("2026-05-01")]
    import loader as L
    expect = apr[L.COL_AMOUNT].sum() / apr["bill_uid"].nunique()
    assert meta["grand"][col] == pytest.approx(expect)


def test_the_portfolio_side_ties_too_and_separates_brands_in_one_mall():
    pf = _pf()
    out, meta = PV.build(pf, feed="portfolio", rows=["Store"],
                         measures=["Sales (₹)"])
    assert meta["grand"]["Sales (₹)"] == pytest.approx(pf["sales"].sum())
    assert set(out["Store"]) == {"MANYAVAR — City Centre", "TURTLE — City Centre"}


def test_text_measure_columns_are_coerced_to_numbers():
    """BILL, QTY, FOOTFALL and Day Target arrive as text with blanks."""
    out, meta = PV.build(_pf(), feed="portfolio", rows=["Brand"],
                         measures=["Bills", "Day Target (₹)"])
    assert meta["grand"]["Bills"] == pytest.approx(3 * len(_pf()))
    assert meta["grand"]["Day Target (₹)"] == pytest.approx(1000 * len(_pf()))


def test_a_sparsely_recorded_measure_is_reported_as_such():
    """A blank column must be explained, not downloaded as a zero.
    See [[feedback-silent-failure-must-speak]]."""
    cov = PV.coverage(_pf(), "portfolio")
    assert cov["Footfall"] == 0.0
    assert cov["Bills"] == 1.0


def test_filters_and_the_date_window_both_narrow_the_source():
    df = _vfl()
    out, meta = PV.build(df, feed="vfl", rows=["Store"], measures=["Sales (₹)"],
                         date_from="2026-05-01", filters={"Store": ["Agartala"]})
    assert list(out["Store"]) == ["Agartala"]
    assert meta["rows_out"] == 62        # Agartala writes two bills a day
    assert "Store: Agartala" in meta["filters"]


def test_asking_for_nothing_is_refused_rather_than_guessed():
    with pytest.raises(ValueError):
        PV.build(_vfl(), feed="vfl", rows=[], measures=["Sales (₹)"])
    with pytest.raises(ValueError):
        PV.build(_vfl(), feed="vfl", rows=["Store"], measures=[])


# --------------------------------------------------------------------------- #
# The workbook
# --------------------------------------------------------------------------- #
def test_every_column_declares_its_kind():
    """A column in none of the buckets prints as a raw float and Excel cannot
    compute on it. See [[feedback-declare-numeric-columns]]."""
    _out, meta = PV.build(_vfl(), feed="vfl", rows=["Store"],
                          measures=["Sales (₹)", "Bills", "Discount %",
                                    "Avg Bill Value / ATV (₹)"])
    k = meta["kinds"]
    assert k["Store"] == "text" and k["Sales (₹)"] == "money"
    assert k["Bills"] == "int" and k["Discount %"] == "pct"


def test_the_workbook_holds_numbers_not_text():
    """Someone opening this has to be able to pivot it again."""
    out, meta = PV.build(_vfl(), feed="vfl", rows=["Store"],
                         measures=["Sales (₹)", "Bills"])
    meta["about"] = {"Feed": "VFL"}
    data = PV.to_excel(out, meta, title="T", subtitle="s")
    import io
    from openpyxl import load_workbook
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Pivot"]
    assert [c.value for c in ws[5]][:3] == ["Store", "Sales (₹)", "Bills"]
    v = ws.cell(row=6, column=2).value
    assert isinstance(v, (int, float)) and not isinstance(v, str)
    assert "##" in ws.cell(row=6, column=2).number_format      # Indian grouping
    assert ws.freeze_panes and ws.auto_filter.ref
    assert "How this was built" in wb.sheetnames


def test_the_workbook_carries_a_bold_total_row():
    out, meta = PV.build(_vfl(), feed="vfl", rows=["Store"], measures=["Sales (₹)"])
    meta["about"] = {}
    import io
    from openpyxl import load_workbook
    ws = load_workbook(io.BytesIO(PV.to_excel(out, meta, title="T")))["Pivot"]
    last = 5 + len(out) + 1
    assert ws.cell(row=last, column=1).value == "TOTAL"
    assert ws.cell(row=last, column=2).font.bold
