"""The Excel-look pivot (pivot_pane) — the layout it sends, the engine request it makes, the file it writes.

The screen itself is JavaScript (assets/pivot_pane.js); these pin the Python half so a layout from the browser always
turns into the same engine request, and the workbook always says what the screen said. Frames come from the
pivotx tests, so CI runs them without the feed — see [[feedback-test-where-it-runs]].
"""
import io

import pandas as pd
import pytest

import pivot_pane as PP
import pivot_ui as PU
import pivotx as P
from tests.test_pivotx import _vfl2

ASOF = pd.Timestamp("2026-05-31")


@pytest.mark.parametrize("feed", ["vfl", "portfolio"])
def test_every_recommended_pivottable_becomes_a_valid_request(feed):
    M = P.measures(feed)
    for name, p in PU.PRESETS[feed].items():
        ui = PP.from_preset(feed, p)
        spec = PP.to_spec(ui, feed, ASOF)
        assert spec.values and all(m in M for m, _h in spec.values), name
        assert spec.layout == "Tabular" and spec.subtotals and spec.grand_rows and spec.grand_cols, name


def test_the_engine_is_always_asked_for_the_whole_tree():
    """Layout, subtotals and grand totals are the browser's: switching them must never change the request."""
    ui = PP.default_ui("vfl")
    a = PP._spec_json(PP.to_spec(ui, "vfl", ASOF))
    for k, v in [("layout", "tabular"), ("subtotals", "off"), ("grand_rows", False), ("grand_cols", False),
                 ("repeat", True), ("blank", True), ("style", "peanuts")]:
        assert PP._spec_json(PP.to_spec({**ui, k: v}, "vfl", ASOF)) == a, k


def test_a_filter_counts_only_while_its_field_is_on_the_pivot():
    """Excel drops a field's filter when the field leaves the pivot; so does this."""
    ui = {**PP.default_ui("vfl"), "rows": ["Store"], "cols": [], "filters_area": ["Brand"],
          "sel": {"Brand": {"mode": "include", "items": ["Manyavar"]}, "Size": {"mode": "include", "items": ["40"]}}}
    spec = PP.to_spec(ui, "vfl", ASOF)
    assert spec.filters == {"Brand": {"mode": "include", "items": ["Manyavar"]}}


def test_sorting():
    ui = PP.default_ui("vfl")
    assert PP.to_spec({**ui, "sort": {"by": "label", "desc": False}}, "vfl", ASOF).sort_by == "Label"
    assert PP.to_spec({**ui, "sort": {"by": "label", "desc": True}}, "vfl", ASOF).sort_by == "Label (Z to A)"
    s = PP.to_spec({**ui, "sort": {"by": "Units", "desc": True}}, "vfl", ASOF)
    assert s.sort_by == "Units" and s.descending
    assert PP.to_spec({**ui, "sort": {"by": "Not a measure"}}, "vfl", ASOF).sort_by == "Label"


def test_a_stale_layout_is_cleaned_not_fatal():
    """A bookmarked link from before a field was renamed opens without it, instead of failing."""
    ui = {**PP.default_ui("vfl"), "rows": ["Region", "Gone field"], "cols": ["Region", "Financial Year"],
          "values": [{"m": "Sales (₹)", "how": "Value"}, {"m": "Retired measure", "how": "Value"},
                     {"m": "ATV / avg bill value (₹)", "how": "% of grand total"}]}
    out = PP._valid(ui, "vfl")
    assert out["rows"] == ["Region"] and out["cols"] == ["Financial Year"]
    assert [v["m"] for v in out["values"]] == ["Sales (₹)"]      # a ratio can't be a share of a total


def test_the_address_bar_round_trip():
    ui = PP.default_ui("vfl")
    assert PP._ui_from_param(PP._ui_param(ui)) == ui
    assert PP._ui_from_param("not-a-layout") is None


def test_custom_dates():
    ui = {**PP.default_ui("vfl"), "period": "Custom", "from": "2026-04-01", "to": "2026-04-30"}
    s = PP.to_spec(ui, "vfl", ASOF)
    assert s.ranges == [["2026-04-01", "2026-04-30"]] and s.period == "Custom"


def test_the_grid_the_browser_gets_matches_the_engine():
    w = P.prepare(_vfl2(), "vfl")
    ui = {**PP.default_ui("vfl"), "rows": ["Region", "Store"], "cols": ["Financial Year"], "filters_area": [],
          "sel": {}, "period": "All data", "values": [{"m": "Sales (₹)", "how": "Value"}], "top": 0,
          "sort": {"by": "label", "desc": False}}
    spec = PP.to_spec(ui, "vfl", ASOF)
    g = P.build(w, "vfl", spec, prepared=True)
    gp = PP.grid_payload(g, "vfl")
    assert gp["rf"] == ["Region", "Store"] and gp["cf"] == ["Financial Year"] and gp["nv"] == 1
    grand = next(r for r in gp["rows"] if r[0] == "grand")
    assert grand[2][-1] == pytest.approx(w["Bill Amount"].sum())
    kinds = {r[0] for r in gp["rows"]}
    assert kinds == {"item", "subtotal", "grand"}                 # the whole tree, every time


def test_the_workbook_is_the_sheet_as_drawn():
    from openpyxl import load_workbook
    sheet = {
        "cols": [140, 110], "style": "light16", "banded_rows": True, "col_headers": True, "summary_below": False,
        "filters": [["Brand", "(All)"]], "freeze": [1, 1],
        "head": [[{"t": "Row Labels"}, {"t": "Sum of Sales (₹)", "al": "right"}]],
        "body": [
            {"kind": "group", "level": 0, "cells": [{"t": "East & NE", "b": True}, {"v": 2500000.0, "k": "money", "nf": {"cat": "auto"}, "b": True}]},
            {"kind": "item", "level": 1, "hidden": True, "cells": [{"t": "Siliguri", "i": 1}, {"v": 2500000.0, "k": "money", "nf": {"cat": "lakh", "dec": 2}}]},
            {"kind": "grand", "level": 0, "cells": [{"t": "Grand Total", "b": True}, {"v": 2500000.0, "k": "money", "nf": {"cat": "auto"}, "b": True}]},
        ],
    }
    wb = load_workbook(io.BytesIO(PP.to_excel(sheet, title="t", about={"Rows": "Region"})))
    ws = wb["PivotTable"]
    assert ws["A1"].value == "Brand" and ws["B1"].value == "(All)"
    assert ws["A3"].value == "Row Labels" and ws["A3"].font.b                       # filters, a blank row, the pivot
    assert ws["B4"].value == 2500000 and ws["B4"].number_format == P._XFMT["money"]
    assert ws["B5"].value == pytest.approx(25.0) and ws["B5"].number_format.endswith('" L"')   # lakhs written as lakhs
    assert ws.row_dimensions[5].outline_level == 1 and ws.row_dimensions[5].hidden            # a collapsed group stays collapsed
    assert ws["A6"].value == "Grand Total" and ws["A6"].font.b
    assert ws.freeze_panes == "B4"
    assert "How this was built" in wb.sheetnames


def test_no_customer_mobile_reaches_the_screen_or_the_file():
    assert "CUSTOMER_MOBILE" in P._SOURCE_DROP and "mobile_clean" in P._SOURCE_DROP
    names = {f for _g, fl in P.GROUPS["vfl"] for f in fl}
    assert not any("mobile" in n.lower() for n in names)
