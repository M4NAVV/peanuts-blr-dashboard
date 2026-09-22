"""Self-serve pivots off the raw feed, as formatted Excel.

Manav, 22 Sep: *"all the different divisions in the company need different
types of datapoints from this raw feed. so we just pivot this raw feed into
the tables they need."*

★ IT BORROWS THE APP'S OWN ARITHMETIC, it does not restate it. The VFL side
aggregates through `loader._agg_base` / `_derive_metric` and names dimensions
with `loader._dim_column`, so a number someone exports here and a number on
the G/D sheet come from one definition. A second copy would drift, and the
export is the copy people would then argue from.

★ NO CUSTOMER MOBILES. The feed carries them and this repo is public; a
builder that let anyone tick "mobile" and take the lot away would be a
different kind of feature. Customers are countable, never listable.
"""
from __future__ import annotations

import io

import pandas as pd

# --------------------------------------------------------------------------- #
#  What each feed can be asked for
# --------------------------------------------------------------------------- #
_PF_TIME = ["Day", "Week", "Month", "Quarter", "Year", "Financial Year",
            "Fiscal Month", "Weekday"]
_PF_CATS = {"Store": "_store", "Brand": "brand", "Location": "location",
            "City": "city", "Region": "region"}

# label -> (column on the frame, how to aggregate). Sparse columns are named as
# such on screen rather than quietly returning zero.
_PF_MEASURES = {
    "Sales (₹)": ("sales", "sum"),
    "Bills": ("_bills", "sum"),
    "Units": ("_units_pf", "sum"),
    "Footfall": ("_footfall", "sum"),
    "Day Target (₹)": ("_target", "sum"),
    "Active stores": ("code", "nunique"),
}
_PF_MONEY = {"Sales (₹)", "Day Target (₹)"}
_PF_RATIOS = {                       # numerator, denominator, is_money
    "Avg Bill Value / ATV (₹)": ("Sales (₹)", "Bills", True),
    "Units per Bill / UPT": ("Units", "Bills", False),
    "Conversion %": ("Bills", "Footfall", False),
}


def fields(feed: str) -> dict:
    """The pickable dimensions and measures for a feed."""
    if feed == "vfl":
        import loader as L
        return {"time": list(L.TIME_DIMS), "cats": list(L.CAT_DIMS),
                "measures": list(L.METRICS),
                "money": {m for m, k in L.METRICS.items() if k in L.MONEY_METRICS}}
    return {"time": list(_PF_TIME), "cats": list(_PF_CATS),
            "measures": list(_PF_MEASURES) + list(_PF_RATIOS),
            "money": set(_PF_MONEY) | {m for m, (_n, _d, mn) in _PF_RATIOS.items() if mn}}


def _pf_prepare(df: pd.DataFrame) -> pd.DataFrame:
    """The portfolio feed with its text measures coerced and a store identity.

    ★ `BILL`, `QTY`, `FOOTFALL` and `Day Target` arrive as TEXT with blanks.
    Left alone they sum to nothing at all; coerced without saying so they would
    report a real zero where the truth is "not recorded".
    """
    import daycal
    w = df.copy()
    w["_store"] = daycal.store_identity(w, ("brand", "location"))
    for out, src in (("_bills", "BILL"), ("_units_pf", "QTY"),
                     ("_footfall", "FOOTFALL"), ("_target", "Day Target")):
        w[out] = (pd.to_numeric(w[src].astype(str).str.replace(",", "", regex=False),
                                errors="coerce")
                  if src in w.columns else pd.NA)
    return w


def coverage(df: pd.DataFrame, feed: str) -> dict:
    """What share of rows actually carries each measure, so a blank column is
    explained on the page rather than downloaded as a zero.
    See [[feedback-silent-failure-must-speak]]."""
    if feed != "portfolio":
        return {}
    w = _pf_prepare(df)
    out = {}
    for label, (col, _how) in _PF_MEASURES.items():
        if col in w.columns and col.startswith("_"):
            out[label] = float(w[col].notna().mean())
    return out


def _pf_dim(work: pd.DataFrame, dim: str, name: str):
    """A label column for a portfolio dimension, ordered the way it reads."""
    if dim in _PF_CATS:
        work[name] = work[_PF_CATS[dim]].astype(str)
        return name, None
    d = pd.to_datetime(work["date"])
    if dim == "Day":
        work[name] = d.dt.strftime("%d %b %Y")
        order = [t.strftime("%d %b %Y") for t in sorted(d.dt.normalize().unique())]
    elif dim == "Week":
        wk = d.dt.to_period("W").dt.start_time
        work[name] = "w/o " + wk.dt.strftime("%d %b %y")
        order = ["w/o " + t.strftime("%d %b %y") for t in sorted(wk.unique())]
    elif dim == "Month":
        work[name] = d.dt.strftime("%b %Y")
        order = [t.strftime("%b %Y")
                 for t in sorted(d.dt.to_period("M").dt.start_time.unique())]
    elif dim == "Quarter":
        work[name] = "Q" + d.dt.quarter.astype(str) + " " + d.dt.year.astype(str)
        order = None
    elif dim == "Year":
        work[name] = d.dt.year.astype(str)
        order = sorted(work[name].unique())
    elif dim == "Financial Year":
        fy = d.dt.year.where(d.dt.month >= 4, d.dt.year - 1)
        work[name] = "FY" + (fy + 1).astype(str).str[2:]
        order = sorted(work[name].unique())
    elif dim == "Fiscal Month":
        work[name] = d.dt.strftime("%b")
        order = ["Apr", "May", "Jun", "Jul", "Aug", "Sep",
                 "Oct", "Nov", "Dec", "Jan", "Feb", "Mar"]
        order = [m for m in order if m in set(work[name])]
    else:                                                   # Weekday
        work[name] = d.dt.day_name()
        order = [w for w in ["Monday", "Tuesday", "Wednesday", "Thursday",
                             "Friday", "Saturday", "Sunday"] if w in set(work[name])]
    return name, order


# --------------------------------------------------------------------------- #
#  Building the table
# --------------------------------------------------------------------------- #
def _base_and_derive(work, group_cols, feed, measures):
    """Base sums per cell, then every asked-for measure derived from them.

    ★ A RATIO IS DERIVED AT EVERY LEVEL IT IS SHOWN, never summed up from the
    level below. ATV over three stores is their total sales over their total
    bills, not the mean of three ATVs. See [[feedback-aggregate-ratios-in-pairs]].
    """
    if feed == "vfl":
        import loader as L
        base = L._agg_base(work, group_cols) if group_cols else \
            L._agg_base(work.assign(_all="All"), ["_all"]).drop(columns=["_all"])
        for label in measures:
            base[label] = L._derive_metric(base, L.METRICS[label])["value"]
        return base

    agg = {}
    for label in measures:
        if label in _PF_MEASURES:
            col, how = _PF_MEASURES[label]
            agg[label] = (col, how)
    for label in measures:                       # ratios need their parts
        if label in _PF_RATIOS:
            for part in _PF_RATIOS[label][:2]:
                agg.setdefault(part, _PF_MEASURES[part])
    if group_cols:
        base = work.groupby(group_cols, dropna=False).agg(**agg).reset_index()
    else:
        base = pd.DataFrame([{k: (work[c].nunique() if h == "nunique"
                                  else work[c].sum())
                              for k, (c, h) in agg.items()}])
    for label in measures:
        if label in _PF_RATIOS:
            num, den, _m = _PF_RATIOS[label]
            d = base[den].where(base[den] != 0)
            base[label] = base[num] / d * (100 if label.endswith("%") else 1)
    return base


def build(df, *, feed, rows, cols=None, measures, date_from=None, date_to=None,
          filters=None, sort_by=None, descending=True, top=None):
    """(frame, meta) — the pivot, and what it took to make it."""
    if not rows:
        raise ValueError("Pick at least one row field.")
    if not measures:
        raise ValueError("Pick at least one measure.")

    work = _pf_prepare(df) if feed == "portfolio" else df.copy()
    n_raw = len(work)
    work["date"] = pd.to_datetime(work["date"], errors="coerce")
    if date_from is not None:
        work = work[work["date"] >= pd.Timestamp(date_from)]
    if date_to is not None:
        work = work[work["date"] <= pd.Timestamp(date_to)]

    applied = []
    for dim, keep in (filters or {}).items():
        if not keep:
            continue
        tmp, _o = _dim(work, dim, "_f")
        work = work[work[tmp].isin(list(keep))].drop(columns=["_f"])
        applied.append(f"{dim}: {', '.join(map(str, keep))}")
    if work.empty:
        return pd.DataFrame(), {"note": "Nothing matches those filters.",
                                "rows_in": n_raw, "rows_out": 0, "filters": applied}

    rcols, orders = [], {}
    for i, dim in enumerate(rows):
        c, o = _dim(work, dim, f"_r{i}")
        rcols.append(c)
        orders[c] = o
    ccol, corder = (None, None)
    if cols:
        ccol, corder = _dim(work, cols, "_c")

    group = rcols + ([ccol] if ccol else [])
    base = _base_and_derive(work, group, feed, measures)
    grand = _base_and_derive(work, [], feed, measures)

    col_spec = {}                      # output column -> (column value, measure)
    if ccol:
        out = base.pivot_table(index=rcols, columns=ccol, values=list(measures),
                               aggfunc="first")
        out = out.reorder_levels([1, 0], axis=1) if len(measures) > 1 else out
        keys = corder or sorted({str(v) for v in base[ccol]})
        if len(measures) > 1:
            want = [(k, m) for k in keys for m in measures if (k, m) in out.columns]
            out = out[want]
            out.columns = [f"{k} · {m}" for k, m in want]
            col_spec = {f"{k} · {m}": (k, m) for k, m in want}
        else:
            m = measures[0]
            want = [(m, k) for k in keys if (m, k) in out.columns]
            out = out[want]
            out.columns = [k for _m, k in want]
            col_spec = {k: (k, m) for _m, k in want}
        out = out.reset_index()
    else:
        out = base[rcols + list(measures)].copy()

    for c in rcols:
        if orders.get(c):
            out[c] = pd.Categorical(out[c], categories=orders[c], ordered=True)
    out = out.sort_values(rcols) if not sort_by else \
        out.sort_values(sort_by, ascending=not descending)
    for c in rcols:
        out[c] = out[c].astype(str)
    if top:
        out = out.head(int(top))
    out.columns = [dict(zip(rcols, rows)).get(c, c) for c in out.columns]

    # ★ EVERY TOTAL IS DERIVED AT ITS OWN LEVEL. A pivoted ratio column cannot
    # be totalled by adding the ratios above it, so the column totals are
    # re-aggregated from the base over exactly the rows that column covers.
    # See [[feedback-aggregate-ratios-in-pairs]].
    total = {rows[0]: "TOTAL"}
    per_col = (_base_and_derive(work, [ccol], feed, measures).set_index(ccol)
               if ccol else None)
    for c in out.columns[len(rows):]:
        if ccol and c in col_spec:
            key, meas = col_spec[c]
            total[c] = (float(per_col.loc[key, meas])
                        if per_col is not None and key in per_col.index else None)
        elif c in grand.columns:
            total[c] = float(grand.iloc[0][c])
        else:
            total[c] = pd.to_numeric(out[c], errors="coerce").sum()
    # ★ EVERY COLUMN DECLARES ITS KIND. A column in none of the buckets prints
    # as a raw float, left-aligned, and Excel cannot compute on it.
    # See [[feedback-declare-numeric-columns]].
    money = fields(feed)["money"]
    counts = {"Bills", "Units", "Active stores", "Unique Customers", "Active Stores"}
    kinds = {c: "text" for c in rows}
    for c in out.columns[len(rows):]:
        meas = col_spec[c][1] if c in col_spec else (measures[0] if ccol else c)
        kinds[c] = ("pct" if str(meas).rstrip().endswith("%") else
                    "money" if meas in money else
                    "int" if meas in counts else "ratio")

    meta = {"rows_in": n_raw, "rows_out": int(len(work)), "filters": applied,
            "cells": int(len(out)), "grand": total, "kinds": kinds,
            "span": (f"{work['date'].min():%d %b %Y} to {work['date'].max():%d %b %Y}")}
    return out.reset_index(drop=True), meta


def _dim(work, dim, name):
    if "store" in work.columns and "net_amount" in work.columns:
        import loader as L
        return L._dim_column(work, dim, name)
    return _pf_dim(work, dim, name)


# --------------------------------------------------------------------------- #
#  The workbook — the house look, and numbers Excel can compute on
# --------------------------------------------------------------------------- #
_MAROON = "7A1F2B"
_CREAM = "FBF7F1"
_RULE = "E8E0D4"
_INK = "2B2B2B"
_MUTED = "6E6052"

# Indian grouping, as a number Excel can still add up. Text would look right
# and be useless to the person who opens it.
_FMT = {"money": '[>=10000000]##\,##\,##\,##0;[>=100000]##\,##\,##0;##,##0',
        "int": "##,##0",
        "ratio": '[>=100000]##\,##\,##0.00;##,##0.00',
        "pct": "0.0%",
        "text": "@"}


def to_excel(frame, meta, *, title, subtitle="", sheet="Pivot") -> bytes:
    """The pivot as a formatted workbook, with a second sheet saying how it
    was built — a file that travels without its maker has to explain itself."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    kinds = meta.get("kinds", {})
    grand = meta.get("grand", {})
    cols = list(frame.columns)

    wb = Workbook()
    ws = wb.active
    ws.title = sheet[:31]

    thin = Side(style="thin", color=_RULE)
    head_fill = PatternFill("solid", fgColor=_MAROON)
    band = PatternFill("solid", fgColor=_CREAM)

    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=15, color=_INK)
    ws["A2"] = subtitle
    ws["A2"].font = Font(size=10, color=_MUTED)
    ws["A3"] = (f"{meta.get('span','')} · {meta.get('rows_out',0):,} source rows"
                + (" · " + " · ".join(meta["filters"]) if meta.get("filters") else ""))
    ws["A3"].font = Font(size=10, color=_MUTED)
    top = 5

    for j, c in enumerate(cols, start=1):
        cell = ws.cell(row=top, column=j, value=str(c))
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = head_fill
        cell.alignment = Alignment(horizontal="left" if kinds.get(c) == "text"
                                   else "right", vertical="center", wrap_text=True)
        cell.border = Border(bottom=thin)

    for i, (_ix, row) in enumerate(frame.iterrows()):
        r = top + 1 + i
        for j, c in enumerate(cols, start=1):
            v = row[c]
            kind = kinds.get(c, "text")
            if kind != "text":
                v = None if pd.isna(v) else float(v)
                if kind == "pct":
                    v = None if v is None else v / 100.0
            cell = ws.cell(row=r, column=j, value=v)
            cell.number_format = _FMT.get(kind, "@")
            cell.alignment = Alignment(horizontal="left" if kind == "text" else "right")
            cell.border = Border(bottom=thin)
            if i % 2:
                cell.fill = band

    if grand:
        r = top + 1 + len(frame)
        for j, c in enumerate(cols, start=1):
            v = grand.get(c, "")
            kind = kinds.get(c, "text")
            if kind != "text" and v not in ("", None) and not pd.isna(v):
                v = float(v) / (100.0 if kind == "pct" else 1.0)
            elif kind != "text":
                v = None
            cell = ws.cell(row=r, column=j, value=v)
            cell.number_format = _FMT.get(kind, "@")
            cell.font = Font(bold=True, color=_INK)
            cell.alignment = Alignment(horizontal="left" if kind == "text" else "right")
            cell.border = Border(top=Side(style="medium", color=_MAROON))

    for j, c in enumerate(cols, start=1):
        longest = max([len(str(c))] + [len(str(x)) for x in frame[c].head(400)])
        ws.column_dimensions[get_column_letter(j)].width = min(max(longest + 4, 11), 42)
    ws.freeze_panes = ws.cell(row=top + 1, column=min(len(cols), 2) + 1)
    ws.auto_filter.ref = (f"A{top}:{get_column_letter(len(cols))}"
                          f"{top + len(frame)}")

    about = wb.create_sheet("How this was built")
    for i, (k, v) in enumerate(meta.get("about", {}).items(), start=1):
        about.cell(row=i, column=1, value=k).font = Font(bold=True, color=_INK)
        about.cell(row=i, column=2, value=str(v)).alignment = Alignment(wrap_text=True)
    about.column_dimensions["A"].width = 22
    about.column_dimensions["B"].width = 90

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
