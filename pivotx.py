"""A real pivot table off the raw feed — rows, columns, values, filters, subtotals, "show values as" —
on screen and as an Excel workbook laid out like one.

Manav, 13 Oct 2026 (the admin's feedback on the 22-Sep Excel Builder): *"we need to build a table format
… where we can select whatever fields we need based on the data that is filled in the Google sheet. There
will be some additional fields like financial division and region division … replicate a pivot table as in
excel where data can be transformed into whatever format we need to study in."*

What it adds over `pivot.py` (which stays, so the old tab can come back in one line):
  • MANY row fields and MANY column fields, nested, with subtotals at every level and grand totals;
  • every useful column of the sheet as a field, plus derived ones: the FINANCIAL calendar (FY, FY half,
    FY quarter, fiscal month), the REGION hierarchy (region → state → city → store), price bracket,
    sale/return, salesperson keyed on the code;
  • values "shown as" % of grand / column / row / parent total, change and % change against the previous
    column (month-on-month, year-on-year), running total;
  • tabular or compact layout, sort by label or by a value, top N inside each group, filters that include
    or exclude.

★ THE ARITHMETIC IS THE APP'S. VFL aggregates through `loader._agg_base` / `_derive_metric`; every
subtotal and total is RE-DERIVED from the rows it covers, never added up from the cells above it, so a
ratio's total is a ratio (April ATV is total sales over total bills, not the mean of store ATVs).
See [[feedback-aggregate-ratios-in-pairs]].
★ NO CUSTOMER MOBILES, as a field or in the source-rows sheet. Customers are countable, never listable.
"""
from __future__ import annotations

import base64
import io
import json
import re
from dataclasses import dataclass, field

import pandas as pd

# =========================================================================== #
#  Fields
# =========================================================================== #
TIME_FIN = ["Financial Year", "FY Half", "FY Quarter", "Fiscal Month"]
TIME_CAL = ["Year", "Quarter", "Month", "Week", "Day", "Weekday", "Day of month"]
_FY_MONTHS = ["Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec", "Jan", "Feb", "Mar"]
_WEEKDAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday", "Sunday"]

# field -> column on the VFL frame (plain label fields)
_VFL_COLS = {
    "Region": "region", "State": "state", "City": "city", "Store": "store", "Store format": "store_format",
    "Brand": "brand", "Men/Women/Child": "Men/Women/Child", "Division": "Division", "Section": "Section",
    "Department": "Department", "Size": "Size", "Colour": "color_name", "Colour (with code)": "CATEGORY2",
    "Style code": "CATEGORY1", "HSN code": "HSN_SAC_CODE", "Tax rate": "Tax rate",
}
_PF_COLS = {"Region": "region", "City": "city", "Location": "location", "Brand": "brand", "Store code": "code"}

GROUPS = {
    "vfl": [
        ("Financial calendar", TIME_FIN),
        ("Calendar", TIME_CAL),
        ("Region", ["Region", "State", "City", "Store", "Store format"]),
        ("Product", ["Brand", "Men/Women/Child", "Division", "Section", "Department", "Size", "Colour",
                     "Colour (with code)", "Style code", "Price bracket", "HSN code", "Tax rate"]),
        ("People and bills", ["Salesperson", "Sale or return"]),
    ],
    "portfolio": [
        ("Financial calendar", TIME_FIN),
        ("Calendar", TIME_CAL),
        ("Region", ["Region", "City", "Location", "Store", "Store code"]),
        ("Brand", ["Brand", "Brand group"]),
    ],
}


def field_names(feed: str) -> list[str]:
    return [f for _g, fs in GROUPS[feed] for f in fs]


def field_group(feed: str) -> dict:
    return {f: g for g, fs in GROUPS[feed] for f in fs}


def _natural(v) -> tuple:
    return tuple(int(t) if t.isdigit() else t.lower() for t in re.split(r"(\d+)", str(v)))


def _time(d: pd.Series, f: str):
    """(labels, order) for a time field. Financial year runs April to March."""
    fy = d.dt.year.where(d.dt.month >= 4, d.dt.year - 1) + 1          # FY27 = Apr 2026 - Mar 2027
    fy_lab = "FY" + (fy % 100).astype(int).astype(str).str.zfill(2)
    fm = (d.dt.month - 4) % 12                                        # 0 = April
    if f == "Financial Year":
        lab, key = fy_lab, fy
    elif f == "FY Half":
        h = (fm // 6 + 1)
        lab, key = fy_lab + " H" + h.astype(str), fy * 10 + h
    elif f == "FY Quarter":
        q = (fm // 3 + 1)
        lab, key = fy_lab + " Q" + q.astype(str), fy * 10 + q
    elif f == "Fiscal Month":
        lab, key = d.dt.strftime("%b"), fm
    elif f == "Year":
        lab, key = d.dt.year.astype(str), d.dt.year
    elif f == "Quarter":
        lab, key = "Q" + d.dt.quarter.astype(str) + " " + d.dt.year.astype(str), d.dt.year * 10 + d.dt.quarter
    elif f == "Month":
        lab, key = d.dt.strftime("%b %Y"), d.dt.year * 100 + d.dt.month
    elif f == "Week":
        wk = d.dt.to_period("W").dt.start_time
        lab, key = "w/o " + wk.dt.strftime("%d %b %y"), wk
    elif f == "Day":
        lab, key = d.dt.strftime("%d %b %Y"), d.dt.normalize()
    elif f == "Weekday":
        lab, key = d.dt.day_name(), d.dt.dayofweek
    else:                                                             # Day of month
        lab, key = d.dt.day.astype(str).str.zfill(2), d.dt.day
    order = (pd.DataFrame({"l": lab, "k": key}).drop_duplicates("l").sort_values("k")["l"].tolist())
    return lab, order


def labels(work: pd.DataFrame, feed: str, f: str):
    """(labels as text, explicit order or None) for field `f` on the prepared frame."""
    if f in TIME_FIN or f in TIME_CAL:
        return _time(work["date"], f)
    if feed == "vfl":
        if f in _VFL_COLS:
            col = _VFL_COLS[f]
            s = work[col] if col in work.columns else pd.Series("(not in feed)", index=work.index)
            return s.fillna("(blank)").astype(str).str.strip().replace("", "(blank)"), None
        if f == "Salesperson":
            name = work.get("Name (Dm Salesperson)", pd.Series("", index=work.index)).fillna("").astype(str).str.strip()
            code = work.get("SALESPERSON_NO", pd.Series("", index=work.index)).fillna("").astype(str).str.strip()
            return (name.where(name != "", "(no name)") + " · " + code.where(code != "", "no code")), None
        if f == "Sale or return":
            return pd.Series("Sale", index=work.index).where(work["_units"].fillna(0) >= 0, "Return"), ["Sale", "Return"]
        if f == "Price bracket":
            import category_pdf as CP
            u = work["_units"].astype(float).where(lambda x: x != 0)
            price = (work["Bill Amount"].astype(float) / u).abs().fillna(0)
            return price.map(CP._bracket), list(CP.BRACKETS)
    else:
        if f in _PF_COLS:
            return work[_PF_COLS[f]].fillna("(blank)").astype(str), None
        if f == "Store":
            return work["_store"], None
        if f == "Brand group":
            return work["is_vfl"].map({True: "VFL brands", False: "Other brands"}).fillna("Other brands"), \
                ["VFL brands", "Other brands"]
    raise KeyError(f"Unknown field: {f}")


# =========================================================================== #
#  Measures
# =========================================================================== #
@dataclass
class Measure:
    label: str
    kind: str                 # money | int | ratio | pct (fraction) | moneyratio
    additive: bool            # can be "% of total"-ed and summed meaningfully
    parts: tuple = ()         # base aggregates it needs
    fn: object = None         # base frame -> series


def _div(a, b):
    return a / b.where(b != 0)


VFL_MEASURES = {m.label: m for m in [
    Measure("Sales (₹)", "money", True, ("sales",), lambda b: b["sales"]),
    Measure("Net sales after discount (₹)", "money", True, ("net_sales",), lambda b: b["net_sales"]),
    Measure("Units", "int", True, ("units",), lambda b: b["units"]),
    Measure("Bills", "int", True, ("bills",), lambda b: b["bills"]),
    Measure("Discount (₹)", "money", True, ("discount",), lambda b: b["discount"]),
    Measure("Tax (₹)", "money", True, ("tax",), lambda b: b["tax"]),
    Measure("ATV / avg bill value (₹)", "moneyratio", False, ("sales", "bills"), lambda b: _div(b["sales"], b["bills"])),
    Measure("UPT / units per bill", "ratio", False, ("units", "bills"), lambda b: _div(b["units"], b["bills"])),
    Measure("ASP / avg selling price (₹)", "moneyratio", False, ("sales", "units"), lambda b: _div(b["sales"], b["units"])),
    Measure("Discount %", "pct", False, ("discount", "sales"), lambda b: _div(b["discount"], b["sales"])),
    Measure("Unique customers", "int", False, ("customers",), lambda b: b["customers"]),
    Measure("Active stores", "int", False, ("stores",), lambda b: b["stores"]),
    Measure("Styles sold", "int", False, ("styles",), lambda b: b["styles"]),
    Measure("Salespeople", "int", False, ("salespeople",), lambda b: b["salespeople"]),
    Measure("Days traded", "int", False, ("days",), lambda b: b["days"]),
]}

PF_MEASURES = {m.label: m for m in [
    Measure("Sales (₹)", "money", True, ("sales",), lambda b: b["sales"]),
    Measure("Bills", "int", True, ("bills",), lambda b: b["bills"]),
    Measure("Units", "int", True, ("units",), lambda b: b["units"]),
    Measure("Footfall", "int", True, ("footfall",), lambda b: b["footfall"]),
    Measure("Day target (₹)", "money", True, ("target",), lambda b: b["target"]),
    Measure("ATV / avg bill value (₹)", "moneyratio", False, ("sales", "bills"), lambda b: _div(b["sales"], b["bills"])),
    Measure("UPT / units per bill", "ratio", False, ("units", "bills"), lambda b: _div(b["units"], b["bills"])),
    Measure("Conversion %", "pct", False, ("bills", "footfall"), lambda b: _div(b["bills"], b["footfall"])),
    Measure("Active stores", "int", False, ("stores",), lambda b: b["stores"]),
    Measure("Days traded", "int", False, ("days",), lambda b: b["days"]),
]}

SHOW_AS = ["Value", "% of grand total", "% of column total", "% of row total", "% of parent row",
           "Change vs previous column", "% change vs previous column", "Running total across"]
SHORT = {"% of grand total": "% of total", "% of column total": "% of column", "% of row total": "% of row",
         "% of parent row": "% of parent", "Change vs previous column": "change", "% change vs previous column": "% change",
         "Running total across": "running total"}
SHOW_ADDITIVE_ONLY = {"% of grand total", "% of column total", "% of row total", "% of parent row", "Running total across"}


def measures(feed: str) -> dict:
    return VFL_MEASURES if feed == "vfl" else PF_MEASURES


def show_as_for(feed: str, m: str) -> list[str]:
    """A ratio or a distinct count cannot be a share of a total, nor run up a total."""
    return SHOW_AS if measures(feed)[m].additive else [s for s in SHOW_AS if s not in SHOW_ADDITIVE_ONLY]


def _num(s):
    return pd.to_numeric(s.astype(str).str.replace(",", "", regex=False).str.replace("%", "", regex=False), errors="coerce")


def prepare(df: pd.DataFrame, feed: str) -> pd.DataFrame:
    """Coerce the text measures once; give the portfolio a store identity. No mobiles leave this frame
    except as the hashed-free count the VFL `customers` aggregate needs."""
    w = df.copy()
    w["date"] = pd.to_datetime(w["date"], errors="coerce")
    if feed == "vfl":
        w["_tax"] = _num(w["Tax Amount"]) if "Tax Amount" in w.columns else 0.0
    else:
        import pivot as PV
        w = PV._pf_prepare(w)
    return w


def _aggregate(work: pd.DataFrame, keys: list[str], feed: str, parts: set) -> pd.DataFrame:
    """Base aggregates per group (one row when keys is empty)."""
    tmp = work if keys else work.assign(_all=1)
    g = keys or ["_all"]
    if feed == "vfl":
        import loader as L
        core = L._agg_base(tmp, g)                                    # the app's own definitions
        extra = {}
        if "tax" in parts:
            extra["tax"] = ("_tax", "sum")
        if "styles" in parts:
            extra["styles"] = ("CATEGORY1", "nunique")
        if "salespeople" in parts:
            extra["salespeople"] = ("SALESPERSON_NO", "nunique")
        if "days" in parts:
            extra["days"] = ("date", "nunique")
        if extra:
            core = core.merge(tmp.groupby(g, dropna=False).agg(**extra).reset_index(), on=g, how="left")
    else:
        spec = {"sales": ("sales", "sum"), "bills": ("_bills", "sum"), "units": ("_units_pf", "sum"),
                "footfall": ("_footfall", "sum"), "target": ("_target", "sum"),
                "stores": ("code", "nunique"), "days": ("date", "nunique")}
        core = tmp.groupby(g, dropna=False).agg(**{k: v for k, v in spec.items() if k in parts}).reset_index()
    return core.drop(columns=["_all"]) if not keys else core


# =========================================================================== #
#  The pivot
# =========================================================================== #
@dataclass
class Spec:
    rows: list = field(default_factory=list)
    cols: list = field(default_factory=list)
    values: list = field(default_factory=list)          # [(measure, show_as)]
    filters: dict = field(default_factory=dict)         # field -> {"mode": "include"|"exclude", "items": [..]}
    date_from: object = None
    date_to: object = None
    ranges: list = field(default_factory=list)         # [[from, to], ...] — e.g. FY to date this year AND last
    period: str = ""                                    # the period's name, for the page and the workbook
    layout: str = "Tabular"                             # Tabular | Compact
    subtotals: bool = True
    grand_rows: bool = True
    grand_cols: bool = True
    sort_by: str = "Label"                              # Label | a measure label
    descending: bool = True
    top: int = 0

    def to_param(self) -> str:
        d = dict(self.__dict__)
        d["date_from"] = str(self.date_from) if self.date_from is not None else None
        d["date_to"] = str(self.date_to) if self.date_to is not None else None
        d["ranges"] = [[str(a), str(b)] for a, b in (self.ranges or [])]
        return base64.urlsafe_b64encode(json.dumps(d, separators=(",", ":")).encode()).decode().rstrip("=")

    @staticmethod
    def from_param(s: str) -> "Spec":
        d = json.loads(base64.urlsafe_b64decode(s + "=" * (-len(s) % 4)).decode())
        d["values"] = [tuple(v) for v in d.get("values", [])]
        return Spec(**d)


MAX_ROWS = 20000


def _in_period(w: pd.DataFrame, spec: Spec) -> pd.DataFrame:
    if spec.ranges:
        keep = pd.Series(False, index=w.index)
        for a, b in spec.ranges:
            keep |= (w["date"] >= pd.Timestamp(a)) & (w["date"] <= pd.Timestamp(b))
        return w[keep]
    if spec.date_from is not None:
        w = w[w["date"] >= pd.Timestamp(spec.date_from)]
    if spec.date_to is not None:
        w = w[w["date"] <= pd.Timestamp(spec.date_to)]
    return w


def periods(asof) -> dict:
    """Named periods. The two-year ones put the SAME dates side by side, so a growth column compares like
    with like in time (a part-year against a whole year would read as a collapse)."""
    a = pd.Timestamp(asof).normalize()
    fy0 = pd.Timestamp(a.year if a.month >= 4 else a.year - 1, 4, 1)
    ly = lambda t: t - pd.DateOffset(years=1)
    m0 = a.replace(day=1)
    lm_end = m0 - pd.Timedelta(days=1)
    return {
        "FY to date · this year and last": [[fy0, a], [ly(fy0), ly(a)]],
        "Month to date · this year and last": [[m0, a], [ly(m0), ly(a)]],
        "Last 30 days · this year and last": [[a - pd.Timedelta(days=29), a], [ly(a - pd.Timedelta(days=29)), ly(a)]],
        "This FY": [[fy0, a]],
        "Last FY": [[ly(fy0), fy0 - pd.Timedelta(days=1)]],
        "This month": [[m0, a]],
        "Last month": [[lm_end.replace(day=1), lm_end]],
        "All data": [],
    }


class TooBig(ValueError):
    pass


def build(df: pd.DataFrame, feed: str, spec: Spec, prepared: bool = False) -> dict:
    """The pivot as a grid: header rows, body rows (items, subtotals, grand total), and its kinds."""
    if not spec.values:
        raise ValueError("Add at least one value.")
    M = measures(feed)
    vals = [(m, s) for m, s in spec.values if m in M]
    w = df if prepared else prepare(df, feed)
    n_in = len(w)
    w = _in_period(w, spec)

    used = list(dict.fromkeys(list(spec.rows) + list(spec.cols) + list(spec.filters)))
    orders, applied = {}, []
    w = w.copy()
    for i, f in enumerate(used):
        lab, order = labels(w, feed, f)
        w[f"_f{i}"] = lab.values
        orders[f] = order
    col_of = {f: f"_f{i}" for i, f in enumerate(used)}
    for f, flt in (spec.filters or {}).items():
        items = [str(x) for x in (flt or {}).get("items", [])]
        if not items:
            continue
        keep = w[col_of[f]].isin(items)
        w = w[keep if flt.get("mode", "include") == "include" else ~keep]
        applied.append(f"{f} {'is' if flt.get('mode', 'include') == 'include' else 'is not'} {', '.join(items)}")
    if w.empty:
        return {"empty": True, "note": "Nothing matches those filters and dates.", "rows_in": n_in, "rows_used": 0,
                "filters": applied}

    R = [col_of[f] for f in spec.rows]
    C = [col_of[f] for f in spec.cols]
    parts = set()
    for m, _s in vals:
        parts |= set(M[m].parts)

    # every (row depth, column depth) a cell can live at, aggregated once each
    rdepths = set(range(len(R) + 1)) if spec.subtotals else {0, len(R)}
    if R and any(h == "% of parent row" for _m, h in vals):
        rdepths.add(len(R) - 1)
    rdepths = sorted(rdepths)
    cdepths = sorted(set(range(len(C) + 1)) if spec.subtotals else {0, len(C)})
    cell = {}
    for d in rdepths:
        for e in cdepths:
            keys = R[:d] + C[:e]
            base = _aggregate(w, keys, feed, parts)
            out = {m: M[m].fn(base).astype(float).tolist() for m, _s in vals}
            kt = base[keys].astype(str).values.tolist() if keys else [[]]
            for i, k in enumerate(kt):
                cell[(tuple(k[:d]), tuple(k[d:]))] = {m: out[m][i] for m in out}

    def v(rk, ck, m):
        x = cell.get((rk, ck), {}).get(m)
        return None if x is None or pd.isna(x) else float(x)

    sort_m = spec.sort_by if spec.sort_by in M else None

    def order_children(level_vals, field_name, prefix):
        o = orders.get(field_name)
        if sort_m:
            def key(x):
                val = v(prefix + (x,), (), sort_m)
                return (val is None, -(val or 0) if spec.descending else (val or 0))
            return sorted(level_vals, key=key)
        if o:
            pos = {x: i for i, x in enumerate(o)}
            s = sorted(level_vals, key=lambda x: (pos.get(x, 10 ** 9), _natural(x)))
            return s[::-1] if spec.sort_by == "Label (Z to A)" else s
        s = sorted(level_vals, key=_natural)
        return s[::-1] if spec.sort_by == "Label (Z to A)" else s

    # ---- the row tree
    full_rows = sorted({rk for (rk, ck) in cell if len(rk) == len(R) and ck == ()})
    tree = {}
    for rk in full_rows:
        for d in range(len(R)):
            tree.setdefault(rk[:d], set()).add(rk[d])
    body, hidden_items = [], 0

    def walk(prefix):
        nonlocal hidden_items
        d = len(prefix)
        if d == len(R):
            body.append(("item", prefix))
            return
        kids = order_children(list(tree.get(prefix, [])), spec.rows[d], prefix)
        if spec.top and d == len(R) - 1 and len(kids) > spec.top:
            tm = sort_m or vals[0][0]
            ranked = sorted(kids, key=lambda x: -(v(prefix + (x,), (), tm) or 0))[:spec.top]
            hidden_items += len(kids) - len(ranked)
            keep = set(ranked)
            kids = [k for k in kids if k in keep]
        for k in kids:
            if spec.layout == "Compact" and d + 1 < len(R) and spec.subtotals:
                body.append(("head", prefix + (k,)))
            walk(prefix + (k,))
            if spec.layout != "Compact" and d + 1 < len(R) and spec.subtotals:
                body.append(("subtotal", prefix + (k,)))

    if R:
        walk(())
    if spec.grand_rows or not R:
        body.append(("grand", ()))
    if len(body) > MAX_ROWS:
        raise TooBig(f"That table would have {len(body):,} rows — more than {MAX_ROWS:,}. "
                     "Add a filter, a shorter period, or fewer row fields.")

    # ---- the column tree
    colent = []
    if C:
        full_cols = sorted({ck for (rk, ck) in cell if rk == () and len(ck) == len(C)})
        ctree = {}
        for ck in full_cols:
            for e in range(len(C)):
                ctree.setdefault(ck[:e], set()).add(ck[e])

        def cwalk(prefix):
            e = len(prefix)
            if e == len(C):
                colent.append(("item", prefix))
                return
            o = orders.get(spec.cols[e])
            kids = list(ctree.get(prefix, []))
            if o:
                pos = {x: i for i, x in enumerate(o)}
                kids.sort(key=lambda x: (pos.get(x, 10 ** 9), _natural(x)))
            else:
                kids.sort(key=_natural)
            for k in kids:
                cwalk(prefix + (k,))
                if e + 1 < len(C) and spec.subtotals:
                    colent.append(("subtotal", prefix + (k,)))
        cwalk(())
        if len(colent) * len(vals) > 400:
            raise TooBig(f"That would make {len(colent) * len(vals):,} columns. Put the long field in Rows, "
                         "or filter it down.")
    if spec.grand_cols or not C:
        colent.append(("grand", ()))

    # previous sibling of each column (for change / running total), at its own level
    prev, run_seq = {}, {}
    sib = {}
    for kind, ck in colent:
        if kind in ("item", "subtotal"):
            sib.setdefault((len(ck), ck[:-1]), []).append(ck)
    for (_l, _p), seq in sib.items():
        for i, ck in enumerate(seq):
            prev[ck] = seq[i - 1] if i else None
            run_seq[ck] = seq[:i + 1]

    def parent_rk(kind, rk):
        return rk[:-1] if rk else None

    def shown(kind, rk, ckind, ck, m, how):
        x = v(rk, ck, m)
        if how == "Value" or x is None and how not in ("Running total across",):
            return x
        if how == "% of grand total":
            t = v((), (), m)
            return x / t if t else None
        if how == "% of column total":
            t = v((), ck, m)
            return x / t if t else None
        if how == "% of row total":
            t = v(rk, (), m)
            return x / t if t else None
        if how == "% of parent row":
            p = parent_rk(kind, rk)
            if p is None:
                return 1.0
            t = v(p, ck, m)
            return x / t if t else None
        if how in ("Change vs previous column", "% change vs previous column"):
            if ckind == "grand" or prev.get(ck) is None:
                return None
            p = v(rk, prev[ck], m)
            if p is None:
                return None
            return (x - p) if how.startswith("Change") else ((x - p) / abs(p) if p else None)
        if how == "Running total across":
            if ckind == "grand":
                return x
            tot = [v(rk, c, m) for c in run_seq.get(ck, [ck])]
            tot = [t for t in tot if t is not None]
            return sum(tot) if tot else None
        return x

    def kind_of(m, how):
        k = M[m].kind
        if how.startswith("%"):
            return "pctchg" if "change" in how else "pct"
        if how.startswith("Change"):
            return {"money": "moneychg", "moneyratio": "moneychg", "int": "intchg"}.get(k, "ratiochg") if k != "pct" else "ppchg"
        return k

    leaf = [(ckind, ck, m, how) for (ckind, ck) in colent for (m, how) in vals]
    kinds = [kind_of(m, how) for (_k, _c, m, how) in leaf]

    rows_out = []
    for kind, rk in body:
        vals_row = [shown(kind, rk, ck_kind, ck, m, how) for (ck_kind, ck, m, how) in leaf]
        rows_out.append({"kind": kind, "key": rk, "depth": len(rk), "values": vals_row})

    # header rows: one per column field, then the value names (when there is more than one value)
    heads = []
    for e in range(len(C)):
        row = []
        for ckind, ck in colent:
            if ckind == "grand":
                txt = "Grand Total" if e == 0 else ""
            elif ckind == "subtotal":
                txt = (ck[e] + " Total") if e == len(ck) - 1 else (ck[e] if e < len(ck) else "")
            else:
                txt = ck[e]
            row.append(txt)
        heads.append(row)
    vnames = [m if how == "Value" else f"{m} {SHORT[how]}" for (m, how) in vals]

    return {
        "empty": False, "spec": spec, "feed": feed, "row_fields": list(spec.rows), "col_fields": list(spec.cols),
        "colent": colent, "vnames": vnames, "nvals": len(vals), "heads": heads, "leaf": leaf, "kinds": kinds,
        "rows": rows_out, "rows_in": n_in, "rows_used": int(len(w)), "filters": applied,
        "hidden_items": hidden_items,
        "span": f"{w['date'].min():%d %b %Y} to {w['date'].max():%d %b %Y}",
    }


# =========================================================================== #
#  Numbers as text (screen) and as Excel formats
# =========================================================================== #
def _indian(n: float, dec: int = 0) -> str:
    neg = n < 0
    n = abs(n)
    s = f"{n:.{dec}f}"
    whole, _, frac = s.partition(".")
    if len(whole) > 3:
        head, tail = whole[:-3], whole[-3:]
        head = ",".join(re.findall(r"\d{1,2}", head[::-1]))[::-1]
        whole = head + "," + tail
    return ("-" if neg else "") + whole + ("." + frac if frac else "")


def fmt(x, kind: str) -> str:
    if x is None or (isinstance(x, float) and pd.isna(x)):
        return ""
    if kind in ("money", "int", "moneyratio"):
        return _indian(x, 0)
    if kind == "ratio":
        return f"{x:,.2f}"
    if kind == "pct":
        return f"{x * 100:.1f}%"
    if kind == "pctchg":
        return f"{x * 100:+.1f}%"
    if kind == "ppchg":
        return f"{x * 100:+.1f} pts"
    if kind in ("moneychg", "intchg"):
        return ("+" if x > 0 else "") + _indian(x, 0)
    if kind == "ratiochg":
        return f"{x:+,.2f}"
    return str(x)


_XFMT = {
    "money": '[>=10000000]##\\,##\\,##\\,##0;[>=100000]##\\,##\\,##0;##,##0',
    "moneyratio": '[>=100000]##\\,##\\,##0;##,##0',
    "int": "##,##0", "ratio": "#,##0.00", "pct": "0.0%", "pctchg": "+0.0%;-0.0%;0.0%",
    "ppchg": '+0.0" pts";-0.0" pts"', "moneychg": "+##,##0;-##,##0;0", "intchg": "+##,##0;-##,##0;0",
    "ratiochg": "+0.00;-0.00;0.00",
}


def _row_labels(g: dict, r: dict) -> list[str]:
    """The label cells of a body row, in the chosen layout."""
    R = g["row_fields"]
    kind, key = r["kind"], r["key"]
    if g["spec"].layout == "Compact":
        if kind == "grand":
            return ["Grand Total"]
        lab = key[-1] if key else ""
        return [lab]
    out = [""] * max(1, len(R))
    if kind == "grand":
        out[0] = "Grand Total"
    elif kind == "subtotal":
        out[len(key) - 1] = f"{key[-1]} Total"
    else:
        for i, k in enumerate(key):
            out[i] = k
    return out


def _blank_repeats(g: dict) -> list[list[str]]:
    """Tabular rows with a label shown only where it changes, as a pivot prints it."""
    out, last = [], None
    for r in g["rows"]:
        labs = _row_labels(g, r)
        if g["spec"].layout != "Compact" and r["kind"] == "item" and last is not None:
            shown_ = list(labs)
            for i in range(len(labs) - 1):
                if last[: i + 1] == list(r["key"][: i + 1]):
                    shown_[i] = ""
                else:
                    break
            labs = shown_
        if r["kind"] == "item":
            last = list(r["key"])
        elif r["kind"] in ("subtotal", "grand"):
            last = None
        out.append(labs)
    return out


# =========================================================================== #
#  On screen
# =========================================================================== #
CSS = """
<style>
.pvx-wrap{max-height:640px;overflow:auto;border:1px solid #e8e0d4;border-radius:6px;background:#fff}
.pvx{border-collapse:separate;border-spacing:0;font:13px/1.35 -apple-system,"Segoe UI",Roboto,sans-serif;color:#2b2b2b;min-width:100%}
.pvx th,.pvx td{padding:6px 10px;border-bottom:1px solid #f0e9df;white-space:nowrap}
.pvx thead th{position:sticky;background:#7a1f2b;color:#fff;font-weight:600;z-index:2;text-align:right;border-bottom:1px solid #8f3542}
.pvx thead th.l{text-align:left}
.pvx thead tr:nth-child(1) th{top:0}.pvx thead tr:nth-child(2) th{top:31px}.pvx thead tr:nth-child(3) th{top:62px}.pvx thead tr:nth-child(4) th{top:93px}
.pvx thead th.grp{text-align:center;border-left:1px solid #8f3542}
.pvx td{text-align:right;font-variant-numeric:tabular-nums}
.pvx td.l{text-align:left;color:#2b2b2b}
.pvx tr.item:nth-child(even) td{background:#fdfbf8}
.pvx tr.subtotal td,.pvx tr.head td{background:#f6efe4;font-weight:600}
.pvx tr.grand td{background:#efe3d0;font-weight:700;border-top:2px solid #7a1f2b;position:sticky;bottom:0}
.pvx td.neg{color:#b3261e}.pvx td.pos{color:#1f6b3a}
.pvx td.ind1{padding-left:26px}.pvx td.ind2{padding-left:42px}.pvx td.ind3{padding-left:58px}.pvx td.ind4{padding-left:74px}
.pvx td.lc{position:sticky;left:0;background:inherit;z-index:1}
</style>
"""


def to_html(g: dict, limit: int = 1500) -> str:
    import html as H
    E = H.escape
    R, C, nv = g["row_fields"], g["col_fields"], g["nvals"]
    compact = g["spec"].layout == "Compact"
    nlab = 1 if compact else max(1, len(R))
    lab_heads = (["Row labels" if R else ""] if compact else (R or [""]))
    nhead = len(C) + (1 if (nv > 1 or not C) else 0)
    h = ['<div class="pvx-wrap"><table class="pvx"><thead>']
    for e in range(len(C)):
        h.append("<tr>")
        if e == 0:
            for j, t in enumerate(lab_heads):
                h.append(f'<th class="l" rowspan="{nhead}">{E(t)}</th>')
        row = g["heads"][e]
        i = 0
        while i < len(row):                         # merge equal neighbours under the same parent
            j = i
            while (j + 1 < len(row) and row[j + 1] == row[i] and row[i] != ""
                   and g["colent"][j + 1][1][:e] == g["colent"][i][1][:e] and g["colent"][j + 1][0] == g["colent"][i][0]):
                j += 1
            span = (j - i + 1) * nv
            h.append(f'<th class="grp" colspan="{span}">{E(row[i])}</th>')
            i = j + 1
        h.append("</tr>")
    if nv > 1 or not C:
        h.append("<tr>")
        if not C:
            for t in lab_heads:
                h.append(f'<th class="l">{E(t)}</th>')
        for _ce in g["colent"]:
            for vn in g["vnames"]:
                h.append(f"<th>{E(vn)}</th>")
        h.append("</tr>")
    h.append("</thead><tbody>")
    labs_all = _blank_repeats(g)
    for n, (r, labs) in enumerate(zip(g["rows"], labs_all)):
        if n >= limit and r["kind"] != "grand":
            continue
        h.append(f'<tr class="{r["kind"]}">')
        for j, t in enumerate(labs[:nlab]):
            ind = f' ind{min(4, r["depth"] - 1)}' if compact and r["kind"] == "item" and r["depth"] > 1 else (
                f' ind{min(4, r["depth"] - 1)}' if compact and r["kind"] == "head" and r["depth"] > 1 else "")
            h.append(f'<td class="l lc{ind}">{E(t)}</td>')
        for x, k in zip(r["values"], g["kinds"]):
            cls = ""
            if x is not None and k.endswith("chg"):
                cls = ' class="neg"' if x < 0 else (' class="pos"' if x > 0 else "")
            h.append(f"<td{cls}>{E(fmt(x, k))}</td>")
        h.append("</tr>")
    h.append("</tbody></table></div>")
    return CSS + "".join(h)


# =========================================================================== #
#  The workbook
# =========================================================================== #
_MAROON, _CREAM, _SUB, _GRAND, _RULE, _INK, _MUTED = "7A1F2B", "FDFBF8", "F6EFE4", "EFE3D0", "E8E0D4", "2B2B2B", "6E6052"
_SOURCE_DROP = {"CUSTOMER_MOBILE", "mobile_clean", "bill_uid", "_provisional", "date_only", "fy_start_year",
                "fy_month_idx", "_units", "Sr No"}


def to_excel(g: dict, *, title: str, about: dict, source: pd.DataFrame | None = None) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    R, C, nv = g["row_fields"], g["col_fields"], g["nvals"]
    compact = g["spec"].layout == "Compact"
    nlab = 1 if compact else max(1, len(R))
    lab_heads = (["Row labels"] if compact else (R or [""]))
    wb = Workbook()
    ws = wb.active
    ws.title = "Pivot"
    thin = Side(style="thin", color=_RULE)
    hfill = PatternFill("solid", fgColor=_MAROON)
    ws["A1"] = title
    ws["A1"].font = Font(bold=True, size=14, color=_INK)
    ws["A2"] = f"{g['span']} · {g['rows_used']:,} source rows" + (" · " + " · ".join(g["filters"]) if g["filters"] else "")
    ws["A2"].font = Font(size=10, color=_MUTED)
    top = 4
    nhead = len(C) + (1 if (nv > 1 or not C) else 0)

    def head(r, c, txt, left=False):
        cell = ws.cell(row=r, column=c, value=txt)
        cell.font = Font(bold=True, color="FFFFFF", size=10)
        cell.fill = hfill
        cell.alignment = Alignment(horizontal="left" if left else "center", vertical="center", wrap_text=True)
        return cell

    for j, t in enumerate(lab_heads, start=1):
        head(top, j, t, left=True)
        if nhead > 1:
            for rr in range(top + 1, top + nhead):
                head(rr, j, None)
            ws.merge_cells(start_row=top, start_column=j, end_row=top + nhead - 1, end_column=j)
    for e in range(len(C)):
        row = g["heads"][e]
        i = 0
        while i < len(row):
            jx = i
            while (jx + 1 < len(row) and row[jx + 1] == row[i] and row[i] != ""
                   and g["colent"][jx + 1][1][:e] == g["colent"][i][1][:e] and g["colent"][jx + 1][0] == g["colent"][i][0]):
                jx += 1
            c0 = nlab + 1 + i * nv
            c1 = nlab + (jx + 1) * nv
            head(top + e, c0, row[i])
            for cc in range(c0 + 1, c1 + 1):
                head(top + e, cc, None)
            if c1 > c0:
                ws.merge_cells(start_row=top + e, start_column=c0, end_row=top + e, end_column=c1)
            i = jx + 1
    if nv > 1 or not C:
        for k, _ce in enumerate(g["colent"]):
            for q, vn in enumerate(g["vnames"]):
                head(top + nhead - 1, nlab + 1 + k * nv + q, vn)

    r0 = top + nhead
    labs_all = _blank_repeats(g)
    for n, (r, labs) in enumerate(zip(g["rows"], labs_all)):
        rr = r0 + n
        style = r["kind"]
        fill = (PatternFill("solid", fgColor=_SUB) if style in ("subtotal", "head") else
                PatternFill("solid", fgColor=_GRAND) if style == "grand" else
                PatternFill("solid", fgColor=_CREAM) if n % 2 else None)
        bold = style in ("subtotal", "head", "grand")
        for j, t in enumerate(labs[:nlab], start=1):
            c = ws.cell(row=rr, column=j, value=t)
            c.font = Font(bold=bold, color=_INK)
            c.alignment = Alignment(horizontal="left", indent=(max(0, r["depth"] - 1) if compact else 0))
            if fill:
                c.fill = fill
            c.border = Border(bottom=thin, top=Side(style="medium", color=_MAROON) if style == "grand" else None)
        for q, (x, k) in enumerate(zip(r["values"], g["kinds"])):
            c = ws.cell(row=rr, column=nlab + 1 + q, value=None if x is None else float(x))
            c.number_format = _XFMT.get(k, "General")
            c.font = Font(bold=bold, color=("B3261E" if (x is not None and k.endswith("chg") and x < 0) else _INK))
            c.alignment = Alignment(horizontal="right")
            if fill:
                c.fill = fill
            c.border = Border(bottom=thin, top=Side(style="medium", color=_MAROON) if style == "grand" else None)
        # collapsible groups, as in an Excel pivot: items sit one level under their subtotal
        if style == "item" and len(R) > 1 and g["spec"].subtotals:
            ws.row_dimensions[rr].outline_level = min(7, len(R) - 1)
    ws.sheet_properties.outlinePr.summaryBelow = not compact

    for j in range(1, nlab + 1):
        longest = max([len(str(lab_heads[j - 1]))] + [len(str(l[j - 1])) for l in labs_all[:2000] if len(l) >= j])
        ws.column_dimensions[get_column_letter(j)].width = min(max(longest + 3, 12), 46)
    for q in range(len(g["leaf"])):
        ws.column_dimensions[get_column_letter(nlab + 1 + q)].width = 15
    ws.freeze_panes = ws.cell(row=r0, column=nlab + 1)

    if source is not None and len(source):
        sh = wb.create_sheet("Source rows")
        src = source.drop(columns=[c for c in source.columns if c in _SOURCE_DROP or c.startswith("_f")], errors="ignore")
        src = src.loc[:, ~src.columns.duplicated()]
        cols = list(src.columns)
        sh.append([str(c) for c in cols])
        for rec in src.itertuples(index=False):
            sh.append([(x.to_pydatetime() if isinstance(x, pd.Timestamp) else (None if (isinstance(x, float) and pd.isna(x)) else x))
                       for x in rec])
        ref = f"A1:{get_column_letter(len(cols))}{len(src) + 1}"
        t = Table(displayName="SourceRows", ref=ref)
        t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
        sh.add_table(t)
        sh.freeze_panes = "A2"
        for j, c in enumerate(cols, start=1):
            sh.column_dimensions[get_column_letter(j)].width = min(max(len(str(c)) + 2, 10), 28)

    ab = wb.create_sheet("How this was built")
    for i, (k, val) in enumerate(about.items(), start=1):
        ab.cell(row=i, column=1, value=k).font = Font(bold=True, color=_INK)
        ab.cell(row=i, column=2, value=str(val)).alignment = Alignment(wrap_text=True, vertical="top")
    ab.column_dimensions["A"].width = 24
    ab.column_dimensions["B"].width = 100
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def source_rows(df: pd.DataFrame, feed: str, spec: Spec) -> pd.DataFrame:
    """The rows behind the pivot (dates and filters applied), without mobiles."""
    w = _in_period(prepare(df, feed), spec)
    for f, flt in (spec.filters or {}).items():
        items = [str(x) for x in (flt or {}).get("items", [])]
        if items:
            lab, _o = labels(w, feed, f)
            keep = lab.isin(items)
            w = w[keep if flt.get("mode", "include") == "include" else ~keep]
    for f in dict.fromkeys(list(spec.rows) + list(spec.cols)):
        w[f], _o = labels(w, feed, f)
    return w.drop(columns=[c for c in w.columns if c in _SOURCE_DROP], errors="ignore")
